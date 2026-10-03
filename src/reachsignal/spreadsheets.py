"""Bounded spreadsheet reads, explicit mapping and portable Excel workbooks."""
from copy import deepcopy
import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import math
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from . import input_format as fmt, portable as io

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 1_000_000_000      # 1,000 MB across the files of one upload
MAX_ROWS = 1_000_000           # per sheet or CSV; one Excel sheet holds 1,048,575 data rows
MAX_CELLS = 10_000_000         # across all sheets of one upload, so the parsed tables fit in memory
MAX_UNZIPPED = 2_500_000_000   # opened size of an .xlsx (its XML is typically 5-10 times the file size)
STYLE_CELLS = 200_000          # exports above this many cells are streamed without per-cell formatting
EXCEL_ROWS = 1_048_575
RESULT_SHEET = re.compile(r"Result \d+ ")  # written by project_workbook; skipped on import
COMMON_LABELS = {"id": "Reference", "name": "Name", "source_id": "Source reference", "note": "Notes", "url": "Source URL", "title": "Source title"}


def normalize(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def label(field):
    return fmt.LABELS.get(field, COMMON_LABELS.get(field, field.replace("_", " ").capitalize()))


def suggest(columns, names):
    """Return only an unambiguous match; never guess the first column."""
    for name in names:
        matches = [c for c in columns if normalize(c) == normalize(name)]
        if len(matches) == 1:
            return matches[0]
    return None


def _frame(rows, title):
    rows = list(rows)
    while rows and all(v is None or v == "" for v in rows[-1]):
        rows.pop()
    if not rows:
        return None
    header = ["" if v is None else str(v).strip() for v in rows[0]]
    while header and header[-1] == "" and all(len(row) <= len(header)-1 or row[len(header)-1] in (None, "") for row in rows[1:]):
        header.pop()
    if not header or any(not c for c in header):
        raise io.DataProblem(f"{title}: put a name in every used column of the first row.")
    if len(set(map(normalize, header))) != len(header):
        raise io.DataProblem(f"{title}: column names must be distinct. Rename the duplicate headings.")
    if len(rows)-1 > MAX_ROWS or len(header) > 80:
        raise io.DataProblem(f"{title}: keep at most {MAX_ROWS:,} rows and 80 columns.")
    data = []
    for index, row in enumerate(rows[1:], 2):
        if any(v not in (None, "") for v in row[len(header):]):
            raise io.DataProblem(f"{title}, row {index}: more values than column headings.")
        values = list(row[:len(header)]) + [None] * max(0, len(header)-len(row))
        if any(v not in (None, "") for v in values):
            data.append(values)
    return pd.DataFrame(data, columns=header, dtype=object)


def load_tables(files):
    """files is [(filename, bytes)]; no file contents enter a shared cache."""
    if not files or sum(len(raw) for _, raw in files) > MAX_BYTES:
        raise io.DataProblem(f"Choose Excel (.xlsx) or UTF-8 CSV files totalling no more than {MAX_BYTES // 1_000_000:,} MB.")
    tables = {}
    cell_count = 0
    for filename, raw in files:
        suffix = Path(filename).suffix.lower()
        parsed = {}
        try:
            if suffix == ".csv":
                text = raw.decode("utf-8-sig")
                try:
                    dialect = csv.Sniffer().sniff(text[:16000], delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(StringIO(text), dialect, strict=True)
                rows = []
                for row in reader:
                    if len(rows) > MAX_ROWS or len(row) > 80:
                        raise io.DataProblem(f"{filename}: keep at most {MAX_ROWS:,} rows and 80 columns.")
                    rows.append(row)
                parsed[Path(filename).stem] = _frame(rows, filename)
            elif suffix == ".xlsx":
                with ZipFile(BytesIO(raw)) as archive:
                    if sum(f.file_size for f in archive.infolist()) > MAX_UNZIPPED or len(archive.infolist()) > 1000:
                        raise io.DataProblem("This workbook is too large when opened. Keep only the sheets and rows you need.")
                formulas = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                cached = None
                try:
                    if len(formulas.worksheets) > 30:
                        raise io.DataProblem("Keep at most 30 sheets in the workbook.")
                    for sheet in formulas.worksheets:
                        if RESULT_SHEET.match(sheet.title):
                            continue  # result sheets of an exported workbook are never model inputs
                        if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > 80:
                            raise io.DataProblem(f"{sheet.title}: keep at most {MAX_ROWS:,} rows and 80 columns, including formatted cells.")
                        rows, pending = [], []
                        for r, row in enumerate(sheet.iter_rows()):
                            if len(rows) > MAX_ROWS:
                                raise io.DataProblem("Too many workbook rows.")
                            values = []
                            for c, cell in enumerate(row):
                                kind = cell.data_type
                                if kind == "f":
                                    pending.append((r, c, cell.coordinate))
                                    values.append(None)
                                elif kind == "e":
                                    raise io.DataProblem(f"{sheet.title}, {cell.coordinate}: fix the Excel error before importing.")
                                else:
                                    values.append(cell.value)
                            rows.append(values)
                        if pending:
                            # Formulas use the result Excel saved with the file; the workbook is read a second time
                            # (values only) just for sheets that contain formulas.
                            if cached is None:
                                cached = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                            wanted = {(r, c): coordinate for r, c, coordinate in pending}
                            last = max(r for r, _, _ in pending)
                            for r, saved in enumerate(cached[sheet.title].iter_rows(max_row=last + 1, values_only=True)):
                                for c, value in enumerate(saved):
                                    coordinate = wanted.pop((r, c), None)
                                    if coordinate is None:
                                        continue
                                    if value is None:
                                        raise io.DataProblem(f"{sheet.title}, {coordinate}: this formula has no saved result. Recalculate and save in Excel, or paste its values before uploading.")
                                    rows[r][c] = value
                            if wanted:
                                coordinate = next(iter(wanted.values()))
                                raise io.DataProblem(f"{sheet.title}, {coordinate}: this formula has no saved result. Recalculate and save in Excel, or paste its values before uploading.")
                        parsed[sheet.title] = _frame(rows, sheet.title)
                finally:
                    formulas.close()
                    if cached is not None:
                        cached.close()
            else:
                raise io.DataProblem("Use Excel .xlsx or CSV. For older .xls files, choose Save As → Excel Workbook (.xlsx).")
        except io.DataProblem:
            raise
        except (ValueError, TypeError, UnicodeError, csv.Error, BadZipFile, KeyError, OSError, SyntaxError) as exc:
            raise io.DataProblem(f"Could not read {filename}. Save a clean .xlsx or UTF-8 CSV with headings in the first row.") from exc
        for name, frame in parsed.items():
            if frame is None or normalize(name) in {"readme", "instructions"}:
                continue
            cell_count += (len(frame)+1) * len(frame.columns)
            if cell_count > MAX_CELLS:
                raise io.DataProblem(f"The combined files contain more than {MAX_CELLS:,} cells. Remove unused sheets and columns, or aggregate small areas.")
            title = name if name not in tables else f"{Path(filename).stem} / {name}"
            if title in tables:
                raise io.DataProblem("Two files have the same sheet names. Rename a file or sheet to distinguish them.")
            tables[title] = frame
    if not tables:
        raise io.DataProblem("No data tables found. Put column headings on the first row and data underneath.")
    return tables


def number(value, comma=False, probability=False):
    if isinstance(value, bool):
        raise io.DataProblem("Use a number, not TRUE or FALSE.")
    if isinstance(value, str):
        value = value.strip()
        percent = value.endswith("%")
        if percent:
            if not probability:
                raise io.DataProblem("A percent sign is only supported for probability columns.")
            value = value[:-1].strip()
        if comma:
            if "." in value:
                raise io.DataProblem("This text uses a dot. Choose dot decimals or remove thousands separators.")
            value = value.replace(",", ".")
        elif "," in value:
            raise io.DataProblem("Choose comma decimals for values such as 1,5. Do not include thousands separators.")
        try:
            result = float(value) / (100 if percent else 1)
        except ValueError as exc:
            raise io.DataProblem("Use a numeric value; leave unknown values blank.") from exc
    else:
        try:
            result = float(value)
        except (ValueError, TypeError) as exc:
            raise io.DataProblem("Use a numeric value; leave unknown values blank.") from exc
    if not math.isfinite(result):
        raise io.DataProblem("Use a finite number; leave unknown values blank.")
    if probability and not 0 <= result <= 1:
        raise io.DataProblem("Probabilities use 0 to 1, or a percent sign (for example 40%).")
    return result


def missing(value):
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, float):
        return value != value
    if isinstance(value, (bool, int, date)):
        return False
    return bool(pd.isna(value))


CHOICES = {
    "lane": {"evidence": ["physical / digital evidence", "physical evidence", "digital evidence"], "customer": ["customer actions", "customer action"], "frontstage": ["visible staff / technology", "visible staff", "front stage", "visible"], "backstage": ["backstage actions", "back stage"], "support": ["support processes", "support process"]},
    "basis": {"observed": [], "assumed": [], "proposed": []},
    "role": {"own": ["owned", "our site", "our store"], "competitor": ["competition", "rival"], "candidate": ["proposed", "new location"]},
}
_CHOICE_LOOKUP = {kind: {normalize(alias): canonical for canonical, aliases in options.items() for alias in [canonical] + aliases}
                  for kind, options in CHOICES.items()}


def cast(value, kind, comma=False):
    empty = missing(value)
    if kind == "basis" and empty:
        return "assumed"
    if kind.startswith("optional_") and empty:
        return "" if kind == "optional_text" else None
    if empty:
        raise io.DataProblem("Choose a column and fill in this value.")
    if kind in {"number", "optional_number", "optional_probability"}:
        return number(value, comma, "probability" in kind)
    text = str(value).strip()
    if kind in CHOICES:
        canonical = _CHOICE_LOOKUP[kind].get(normalize(text))
        if canonical is None:
            raise io.DataProblem("Use one of: " + ", ".join(CHOICES[kind]) + ".")
        return canonical
    return text


def mapped_rows(frame, mapping, fields, comma=False, table="Table"):
    used = [c for c in mapping.values() if c is not None]
    if len(used) != len(set(used)):
        raise io.DataProblem(f"{table}: each source column can fill only one role. Check the column choices.")
    columns = {}
    for field, (title, kind, _) in fields.items():
        column = mapping.get(field)
        values = frame[column].tolist() if column else [None] * len(frame)
        out = []
        for index, value in enumerate(values, 1):
            try:
                out.append(cast(value, kind, comma))
            except io.DataProblem as exc:
                raise io.DataProblem(f"{table}, data row {index}, {title}: {exc}") from exc
        columns[field] = out
    return _records(columns, len(frame))


def _records(columns, count):
    names = list(columns)
    return [dict(zip(names, values)) for values in zip(*columns.values())] if names else [{} for _ in range(count)]


def schema_rows(model, table, frame, mapping, comma=False):
    """Convert mapped workbook cells without supplying made-up numeric defaults."""
    specs = model.SCHEMA["properties"][table]["items"]["properties"]
    used = [v for v in mapping.values() if v]
    if len(used) != len(set(used)):
        raise io.DataProblem(f"{fmt.TABLE_NAMES[table]}: choose a different source column for each role.")
    columns = {}
    for field, spec in specs.items():
        actual = spec.get("anyOf", [spec])[0]
        nullable = any(s.get("type") == "null" for s in spec.get("anyOf", []))
        source = mapping.get(field)
        values = frame[source].tolist() if source is not None and source in frame.columns else [None] * len(frame)
        out = []
        for index, value in enumerate(values, 1):
            try:
                if missing(value):
                    if nullable:
                        value = None
                    elif actual.get("type") == "string" and "minLength" not in actual and "pattern" not in actual and "format" not in actual:
                        value = ""
                    else:
                        raise io.DataProblem("Choose a column and supply a value.")
                elif actual.get("type") in {"number", "integer"}:
                    value = number(value, comma, field == "probability")
                    if actual["type"] == "integer":
                        if not value.is_integer():
                            raise io.DataProblem("Use a whole number.")
                        value = int(value)
                elif actual.get("format") == "date" and isinstance(value, (date, datetime)):
                    value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                elif "enum" in actual:
                    matched = suggest(actual["enum"], [str(value)])
                    if matched is None:
                        raise io.DataProblem("Use one of: " + ", ".join(actual["enum"]) + ".")
                    value = matched
                else:
                    value = str(value).strip()
                out.append(value)
            except io.DataProblem as exc:
                raise io.DataProblem(f"{fmt.TABLE_NAMES[table]}, data row {index}, {label(field)}: {exc}") from exc
        columns[field] = out
    return _records(columns, len(frame))


def complete_project(model, brief, tables):
    d = {"schema_version": "1.0", "brief": brief}
    for name in model.TITLES:
        d[name] = deepcopy(tables.get(name, []))
        minimum = model.SCHEMA["properties"][name].get("minItems", 0)
        if len(d[name]) < minimum:
            raise io.DataProblem(f"{fmt.TABLE_NAMES[name]} needs at least {minimum} row(s). Choose its sheet and check the columns.")
    return model.validate(d)


def _workbook(tables, instructions):
    data = {"Read me": pd.DataFrame({"Instructions": instructions}), **tables}
    for title, frame in data.items():
        if len(frame) > EXCEL_ROWS:
            raise io.DataProblem(f"{title} has {len(frame):,} rows, more than one Excel sheet holds. Use the ZIP export.")
    if sum(len(frame) * max(1, len(frame.columns)) for frame in data.values()) > STYLE_CELLS:
        return _streamed_workbook(data)
    book = Workbook()
    book.remove(book.active)
    for title, frame in data.items():
        safe_title = re.sub(r"[\\/*?:\[\]]", " ", title)[:31]
        sheet = book.create_sheet(safe_title)
        for row in [list(frame.columns)] + frame.astype(object).where(pd.notna(frame), None).values.tolist():
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Preserve literal text without creating an Excel formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="343229")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, column in enumerate(frame.columns, 1):
            values = [str(column)] + [str(v) for v in frame[column].head(30) if v is not None]
            sheet.column_dimensions[get_column_letter(i)].width = min(60, max(18, max(map(len, values), default=18)+2))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _streamed_workbook(data):
    """Large exports: openpyxl's write-only mode keeps memory flat. Headers keep their style; text that Excel would
    read as a formula or an error code is still stored as literal text."""
    book = Workbook(write_only=True)
    fill, font = PatternFill("solid", fgColor="343229"), Font(color="FFFFFF", bold=True)

    def text(sheet, value):
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell

    for title, frame in data.items():
        sheet = book.create_sheet(re.sub(r"[\\/*?:\[\]]", " ", title)[:31])
        for i, column in enumerate(frame.columns, 1):
            values = [str(column)] + [str(v) for v in frame[column].head(30) if v is not None]
            sheet.column_dimensions[get_column_letter(i)].width = min(60, max(18, max(map(len, values), default=18)+2))
        sheet.freeze_panes = "A2"
        header = []
        for column in frame.columns:
            cell = text(sheet, str(column))
            cell.fill, cell.font = fill, font
            header.append(cell)
        sheet.append(header)
        for row in frame.astype(object).where(pd.notna(frame), None).itertuples(index=False, name=None):
            sheet.append([text(sheet, v) if isinstance(v, str) and (v.startswith("=") or v in ERROR_CODES) else v for v in row])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def simple_template():
    tables = {}
    for key, spec in fmt.QUICK.items():
        tables[spec["title"]] = pd.DataFrame(fmt.EXAMPLES[key], columns=list(spec["fields"])).rename(columns={k: v[0] for k, v in spec["fields"].items()})
    return _workbook(tables, ["FICTIONAL EXAMPLE — replace every example row with your own data before using results.", fmt.INTRO, fmt.HELP, "Keep the first row as column headings. Blank numeric cells mean unknown, not zero.", "Upload using Simple business tables. Confirm the sheet and column suggestions before importing."])


def csv_template(key):
    spec = fmt.QUICK[key]
    frame = pd.DataFrame(fmt.EXAMPLES[key], columns=list(spec["fields"])).rename(columns={k: v[0] for k, v in spec["fields"].items()})
    return io.csv_bytes(frame)


def project_workbook(model, data, results=None):
    tables = {"Case": pd.DataFrame({"Case brief": [data["brief"]]})}
    for name in model.TITLES:
        columns = list(model.SCHEMA["properties"][name]["items"]["properties"])
        tables[fmt.TABLE_NAMES[name]] = pd.DataFrame(data[name], columns=columns).rename(columns={c: label(c) for c in columns})
    for i, (name, frame) in enumerate((results or {}).items(), 1):
        if name not in model.TITLES:
            tables[f"Result {i} {name}"[:31]] = frame
    return _workbook(tables, ["Signal project workbook. Imported workbooks always need a new review; review signatures are not imported from Excel.", "If this is the fictional example, replace example rows and the case brief before using it for your business.", "Import using Complete project workbook. Sheets and columns are suggested automatically; confirm or correct them.", "Reference columns link related tables. Preserve references when renaming a record. Empty numeric cells remain unknown.", "Result sheets are exports only; the importer does not use them as model inputs."])
